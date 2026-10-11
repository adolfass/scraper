#!/usr/bin/env python3
"""Build the final file for Итоговый ИП.xlsx in Результат.xlsx format.

Input : /Users/mac/Downloads/Итоговый ИП.xlsx  (Brand, Code, Name, Qty, Price)
Data  : ip_probe_log.jsonl / ip_search_rows.jsonl  (new probe)
        probe_log.jsonl / search_rows.jsonl         (previous scan, merged)
        certs.jsonl                                 (detail cards)
        tnved_tree.json
Output: Итоговый ИП_Результат.xlsx  (same 12 columns as Результат.xlsx)

Brand gate: a link is attached only when BrandName matches the cert's
manufacturer/applicant/trademark (summary fields or detail card).
"""
import json, os, re
from collections import defaultdict, Counter

import openpyxl

import pipeline
import build_output as B

ROOT = pipeline.ROOT
SRC = "/Users/mac/Downloads/Итоговый ИП.xlsx"
OUT = os.environ.get("OUT_FILE") or os.path.join(ROOT, "Итоговый ИП_Результат.xlsx")
REF = os.path.join(ROOT, "Результат.xlsx")   # carry-over source + header template

IP_PROBE = os.path.join(ROOT, "ip_probe_log.jsonl")
IP_ROWS = os.path.join(ROOT, "ip_search_rows.jsonl")
DECL_PROBE = os.path.join(ROOT, "ip_decl_probe_log.jsonl")
DECL_ROWS = os.path.join(ROOT, "ip_decl_rows.jsonl")
TNVED_TREE = os.path.join(ROOT, "tnved_tree.json")
DECL_URL = "https://pub.fsa.gov.ru/rds/declaration/view/{id}/product"

# Exact-article override: attach a candidate cert whose identification text
# carries the article VERBATIM even if the manufacturer gate fails (the cert is
# issued to the actual component maker, not the retail brand). Opt-in via env.
EXACT_OVERRIDE = os.environ.get("EXACT_OVERRIDE", "0") == "1"

HEADER = ["ProdNum", "Code", "BrandName", "ProdName", "PriceOut", "QTY",
          "Commentary", "BatchQty", "CertNumber", "CertUrl", "TNVED", "CertType"]
SUMMARY_FIELDS = ("manufacterName", "applicantName",
                  "productIdentificationTrademark",
                  "productIdentificationName", "productFullName")


def load_input():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Лист1"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def build_candidates():
    """code -> set of cert ids, from new + previous probes.

    Match rule follows build_output.RULE (env RULE):
      t   = exact list item only
      any = every pair (substring "s" included)
      mix = "t" everywhere + "s" only for complex codes (default)

    When SUBSTR_TIER2=1, substring-only ("s") ids are also collected in a
    second map so the chooser can require the brand to match the MANUFACTURER
    (tier 2) for them, not merely a mention in the article annex.
    """
    cand = defaultdict(set)
    subs = defaultdict(set)
    complex_re = re.compile(r"[;,\s]")
    tier2_sub = os.environ.get("SUBSTR_TIER2", "0") == "1"
    for path in (pipeline.PROBE_LOG, IP_PROBE):
        if not os.path.exists(path):
            continue
        for rec in pipeline.load_jsonl(path, key="p").values():
            for m in rec.get("matches", []):
                code, rid = m[0], m[1]
                rule = m[2] if len(m) > 2 else "t"
                if B.RULE == "t" and rule != "t":
                    continue
                if B.RULE == "mix" and rule != "t" and not complex_re.search(code):
                    continue
                cand[code].add(rid)
                if tier2_sub and rule != "t":
                    subs[code].add(rid)
    return cand, subs


def build_decl_candidates():
    """code -> set of declaration ids, mirroring build_candidates() rules."""
    cand = defaultdict(set)
    complex_re = re.compile(r"[;,\s]")
    if not os.path.exists(DECL_PROBE):
        return cand
    for rec in pipeline.load_jsonl(DECL_PROBE, key="p").values():
        for m in rec.get("matches", []):
            code, rid = m[0], m[1]
            rule = m[2] if len(m) > 2 else "t"
            if B.RULE == "t" and rule != "t":
                continue
            if B.RULE == "mix" and rule != "t" and not complex_re.search(code):
                continue
            cand[code].add(rid)
    return cand


def main():
    rows = load_input()
    print(f"input rows: {len(rows)}", flush=True)

    cand, subs = build_candidates()
    print(f"codes with candidates: {len(cand)}", flush=True)

    rows_map = pipeline.load_jsonl(pipeline.SEARCH_ROWS, key="id")
    rows_map.update(pipeline.load_jsonl(IP_ROWS, key="id"))
    certs = {k: v["data"] for k, v in
             pipeline.load_jsonl(pipeline.CERTS_LOG, key="id").items()
             if v.get("data") is not None}
    print(f"summaries: {len(rows_map)}, details: {len(certs)}", flush=True)

    id2code = {}
    if os.path.exists(TNVED_TREE):
        tree = json.load(open(TNVED_TREE, encoding="utf-8"))
        id2code = {int(k): v.get("code") for k, v in tree.items()
                   if v.get("code") and str(v.get("code")).isdigit()}

    # brand haystacks: strong from detail card, fallback from search summary
    strong = {}
    weak = {}
    for cid, det in certs.items():
        s, w = B.cert_haystacks(det)
        strong[cid] = s
        weak[cid] = w
    for cid, s in rows_map.items():
        if cid in strong:
            continue
        parts = [s.get(k) for k in SUMMARY_FIELDS if s.get(k)]
        strong[cid] = B.norm_text(" ".join(parts))
        weak[cid] = ""

    # declarations registry (rds) - fills rows the certificate registry misses
    decl_rows = pipeline.load_jsonl(DECL_ROWS, key="id") if os.path.exists(DECL_ROWS) else {}
    decl_cand = build_decl_candidates()
    print(f"declaration summaries: {len(decl_rows)}, "
          f"codes with decl candidates: {len(decl_cand)}", flush=True)
    dstrong, dweak = {}, {}
    for rid, s in decl_rows.items():
        parts = [s.get(k) for k in SUMMARY_FIELDS if s.get(k)]
        dstrong[rid] = B.norm_text(" ".join(parts))
        dweak[rid] = B.norm_text(s.get("article") or "")

    def dtier(rid, brand):
        s = dstrong.get(rid)
        if s is None:
            return 0
        if B.brand_matches(brand, s):
            return 2
        if B.BRAND_WEAK and B.brand_matches(brand, dweak.get(rid) or ""):
            return 1
        return 0

    def choose_decl(rids, brand):
        def k(rid):
            r = decl_rows.get(rid, {})
            return (1 if r.get("idStatus") == 6 else 0, str(r.get("date") or ""))
        ordered = sorted(rids, key=k, reverse=True)
        for lvl in ((2, 1) if B.BRAND_WEAK else (2,)):
            for rid in ordered:
                if dtier(rid, brand) >= lvl:
                    return rid
        return None

    def tier(rid, brand):
        s = strong.get(rid)
        if s is None:
            return 0
        if B.brand_matches(brand, s):
            return 2
        if B.BRAND_WEAK and B.brand_matches(brand, weak.get(rid) or ""):
            return 1
        return 0

    def exact_ok(rid, code):
        w = weak.get(rid)
        return bool(w) and f" {code} " in w

    def choose_g(rids, brand, subs=None, code=None):
        subs = subs or set()
        def k(rid):
            r = rows_map.get(rid, {})
            return (1 if r.get("idStatus") == 6 else 0, str(r.get("date") or ""))
        ordered = sorted(rids, key=k, reverse=True)
        for lvl in ((2, 1) if B.BRAND_WEAK else (2,)):
            for rid in ordered:
                if lvl < 2 and rid in subs:
                    continue          # substring match must reach the manufacturer
                if tier(rid, brand) >= lvl:
                    return rid
        if EXACT_OVERRIDE and code:
            for rid in ordered:
                if exact_ok(rid, code):
                    return rid
        return None

    # carry-over from the reference output, keyed by (code, brand)
    ref_map = {}
    if os.path.exists(REF):
        wb = openpyxl.load_workbook(REF, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        h = list(next(it)); idx = {k: i for i, k in enumerate(h)}
        for r in it:
            code = str(r[idx["Code"]] or "").strip().lower()
            brand = str(r[idx["BrandName"]] or "").strip().lower()
            num = str(r[idx["CertNumber"]] or "").strip()
            if code and brand and num:
                ref_map.setdefault((code, brand), (
                    num, str(r[idx["CertUrl"]] or "").strip(),
                    str(r[idx["TNVED"]] or "").strip(),
                    str(r[idx["CertType"]] or "").strip()))
        wb.close()
    print(f"reference carry-over pairs: {len(ref_map)}", flush=True)

    st = Counter()
    chosen = {}
    decl_choice = {}
    for r in rows:
        brand = str(r[0] or "").strip().lower() if r[0] is not None else ""
        code = str(r[1] or "").strip().lower() if r[1] is not None else ""
        if brand in B.EXCLUDE_BRANDS:
            st["excluded_brand"] += 1
            continue
        if not code:
            st["no_code"] += 1
            continue
        key = (code, brand)
        if key not in chosen:
            rids = cand.get(code)
            chosen[key] = choose_g(rids, brand,
                                   subs.get(code), code) if rids else None
            if rids and chosen[key] is None:
                st["brand_blocked"] += 1
    chosen_ids = {c for c in chosen.values() if c}
    missing = sorted(c for c in chosen_ids if c not in certs)
    print(f"chosen certs: {len(chosen_ids)}, missing details: {len(missing)}",
          flush=True)
    if os.environ.get("FETCH", "1") == "1":
        for i, cid in enumerate(missing, 1):
            try:
                data = pipeline.get_cert(cid)
            except SystemExit:
                # persistent server timeouts: skip this card, reset breaker,
                # keep going (TNVED will just be empty for this cert)
                print(f"  cert {cid}: server gave up, skipping", flush=True)
                pipeline._consec_fail[0] = 0
                pipeline._backoff[0] = pipeline.BACKOFF_BASE
                continue
            except Exception as e:
                print(f"  cert {cid} failed: {e}", flush=True)
                continue
            pipeline._append(pipeline.CERTS_LOG, {"id": cid, "data": data})
            s, w = B.cert_haystacks(data)
            strong[cid] = s
            weak[cid] = w
            if i % 25 == 0:
                print(f"  details {i}/{len(missing)}", flush=True)
        print("details complete", flush=True)
        # re-choose now that details are available
        for key in chosen:
            code, brand = key
            rids = cand.get(code)
            chosen[key] = (choose_g(rids, brand, subs.get(code), code)
                           if rids else None)

    # build output
    out = []
    for r in rows:
        brand = str(r[0] or "").strip() if r[0] is not None else ""
        brand_l = brand.lower()
        code = str(r[1] or "").strip().lower() if r[1] is not None else ""
        if brand_l in B.EXCLUDE_BRANDS:
            continue
        vals = ("", "", "", "")
        cid = chosen.get((code, brand_l))
        if cid:
            m = rows_map.get(cid, {})
            tnved = ""
            detail = certs.get(cid)
            if detail is not None:
                ident = B.ident_for_code(detail, code)
                if ident:
                    ids = ident.get("idTnveds") or []
                    ids = [t.get("id") if isinstance(t, dict) else t for t in ids]
                    tnved = ";".join(dict.fromkeys(
                        id2code[t] for t in ids if t in id2code))
            vals = (m.get("number") or "", B.CERT_URL.format(id=cid),
                    B.cap_tnved(tnved), B.cert_type_of(m))
            st["matched"] += 1
            if vals[2]:
                st["matched_tnved"] += 1
        elif (code, brand_l) in ref_map:
            num, url, tnved, ctype = ref_map[(code, brand_l)]
            vals = (num, url, B.cap_tnved(tnved), ctype)
            st["matched"] += 1
            st["from_reference"] += 1
            if vals[2]:
                st["matched_tnved"] += 1
        else:
            did = None
            key = (code, brand_l)
            if code and code in decl_cand:
                if key not in decl_choice:
                    decl_choice[key] = choose_decl(decl_cand[code], brand_l)
                did = decl_choice[key]
            if did:
                m = decl_rows.get(did, {})
                vals = (m.get("number") or "", DECL_URL.format(id=did),
                        B.cap_tnved(m.get("tnved") or ""), "Декларация")
                st["matched"] += 1
                st["from_declaration"] += 1
                if vals[2]:
                    st["matched_tnved"] += 1
            else:
                st["unmatched"] += 1
        # canonical order: ProdNum, Code, BrandName, ProdName, PriceOut, QTY,
        #                  Commentary, BatchQty, + 4 cert columns
        out.append(("", str(r[1] or ""), str(r[0] or ""), str(r[2] or ""),
                    r[4] if r[4] is not None else "",
                    r[3] if r[3] is not None else "",
                    "", "") + vals)

    print("stats:", dict(st), flush=True)
    owb = openpyxl.Workbook(write_only=True)
    ows = owb.create_sheet("PRICE")
    ows.append(HEADER)
    for row in out:
        ows.append(list(row))
    owb.save(OUT)
    print(f"written {OUT} ({len(out)} rows)", flush=True)


if __name__ == "__main__":
    main()
