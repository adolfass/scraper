#!/usr/bin/env python3
"""Fill still-empty rows of Итоговый ИП_Результат.xlsx from harvested annexes.

Reads /tmp/annex/fills.json (code -> chosen cert from annex_match.py), resolves
cert number / type (search index) and TNVED (detail card idents), and writes an
updated workbook.  Existing non-empty rows are never touched.

Env:
  ANNEX_OUT  output path (default Итоговый ИП_Результат_annex.xlsx)
  ANNEX_APPLY_ONLY_GATED=1  only attach brand-gated matches
"""
import json
import os
import sys
from collections import Counter

import openpyxl

import pipeline
import build_output as B

ROOT = pipeline.ROOT
IN_FILE = "Итоговый ИП_Результат.xlsx"
OUT_FILE = os.environ.get("ANNEX_OUT") or os.path.join(ROOT, "Итоговый ИП_Результат_annex.xlsx")
IDX = "/tmp/annex/certs_index.json"
FILLS = "/tmp/annex/fills.json"
TNVED_TREE = os.path.join(ROOT, "tnved_tree.json")
ONLY_GATED = os.environ.get("ANNEX_APPLY_ONLY_GATED", "0") == "1"


def load_tnved():
    id2code = {}
    if os.path.exists(TNVED_TREE):
        tree = json.load(open(TNVED_TREE, encoding="utf-8"))
        id2code = {int(k): v.get("code") for k, v in tree.items()
                   if v.get("code") and str(v.get("code")).isdigit()}
    return id2code


def union_tnved(detail, id2code):
    idents = (detail.get("product") or {}).get("identifications") or []
    out = []
    for ident in idents:
        ids = ident.get("idTnveds") or []
        ids = [t.get("id") if isinstance(t, dict) else t for t in ids]
        for t in ids:
            if t in id2code:
                out.append(id2code[t])
    return ";".join(dict.fromkeys(out))


def main():
    fills = json.load(open(FILLS, encoding="utf-8"))
    idx = {}
    if os.path.exists(IDX):
        idx = json.load(open(IDX, encoding="utf-8")).get("certs", {})
    id2code = load_tnved()

    cards = {k: v["data"] for k, v in
             pipeline.load_jsonl(pipeline.CERTS_LOG, key="id").items()
             if v.get("data") is not None}
    print(f"fills={len(fills)} index={len(idx)} cards={len(cards)}", flush=True)

    # resolve chosen cert -> (number, url, tnved, ctype)
    resolved = {}
    missing = []
    for code, v in fills.items():
        if ONLY_GATED and not v.get("brand_ok"):
            continue
        cid = str(v["cid"])
        if cid not in cards:
            missing.append(cid)
            resolved[code] = {"cid": cid, "need": True}
        else:
            resolved[code] = {"cid": cid, "need": False}
    missing = sorted(set(missing))
    print(f"chosen certs missing detail: {len(missing)}", flush=True)

    if missing and os.environ.get("ANNEX_FETCH", "1") == "1":
        for i, cid in enumerate(missing, 1):
            try:
                data = pipeline.get_cert(cid)
            except Exception as e:
                print(f"  cert {cid} failed: {str(e)[:80]}", flush=True)
                continue
            cards[cid] = data
            pipeline._append(pipeline.CERTS_LOG, {"id": cid, "data": data})
            if i % 25 == 0:
                print(f"  details {i}/{len(missing)}", flush=True)
        print("detail fetch complete", flush=True)

    # compute values
    vals_by_code = {}
    for code, r in resolved.items():
        cid = r["cid"]
        s = idx.get(cid, {})
        detail = cards.get(cid)
        number = s.get("number")
        ctype = None
        if s:
            ctype = B.cert_type_of(s)
        if detail is not None:
            m = detail
            number = number or m.get("number")
            if ctype is None:
                ctype = B.cert_type_of(m)
        url = B.CERT_URL.format(id=cid)
        tnved = union_tnved(detail, id2code) if detail is not None else ""
        gated = fills[code]["brand_ok"]
        vals_by_code[code] = (number or "", url, B.cap_tnved(tnved), ctype or "Сертификат", gated)

    wb = openpyxl.load_workbook(IN_FILE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    idx_h = {h: i for i, h in enumerate(header)}
    ci = idx_h["Code"]
    ni = idx_h["CertNumber"]
    st = Counter()
    rows = []
    for r in it:
        row = list(r)
        code = str(row[ci] or "").strip().lower()
        empty = not (row[ni] and str(row[ni]).strip())
        if empty and code in vals_by_code:
            number, url, tnved, ctype, gated = vals_by_code[code]
            row[idx_h["CertNumber"]] = number
            row[idx_h["CertUrl"]] = url
            row[idx_h["TNVED"]] = tnved
            row[idx_h["CertType"]] = ctype
            st["filled"] += 1
            st["gated" if gated else "no_gate"] += 1
            if tnved:
                st["with_tnved"] += 1
        # normalise to the exact header width (openpyxl trims trailing Nones)
        row = ["" if v is None else v for v in row]
        if len(row) < len(header):
            row += [""] * (len(header) - len(row))
        rows.append(row)
    wb.close()

    owb = openpyxl.Workbook(write_only=True)
    ows = owb.create_sheet("PRICE")
    ows.append(header)
    for r in rows:
        ows.append(r)
    owb.save(OUT_FILE)
    print("stats:", dict(st), flush=True)
    print(f"written {OUT_FILE} ({len(rows)} rows)", flush=True)


if __name__ == "__main__":
    main()
