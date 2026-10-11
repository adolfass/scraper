#!/usr/bin/env python3
"""Fallback: attach DECLARATION-registry group certificates for brands whose
certificate was not found in the certificate registry.

Same brand-group logic as brand_group.py, but queries
POST /api/v1/rds/common/declarations/get and attaches
https://pub.fsa.gov.ru/rds/declaration/view/{id}/product with
CertType = "Декларация".

Env:
  DG_IN / DG_OUT / DG_CACHE / DG_NLIVE / DG_NTOP
"""
import json
import os
import re
import time
from collections import defaultdict, Counter

import openpyxl

import fsa_api
import build_output as B

WORK = "/tmp/annex"
TNVED_TREE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tnved_tree.json")
IN_FILE = os.environ.get("DG_IN", "/tmp/Итоговый_ИП_final2.xlsx")
OUT_FILE = os.environ.get("DG_OUT", "/tmp/Итоговый_ИП_decl.xlsx")
CACHE = os.environ.get("DG_CACHE", f"{WORK}/decl_search.json")
CARDS = os.environ.get("DG_CARDS", f"{WORK}/decl_cards.json")
NLIVE = int(os.environ.get("DG_NLIVE", "60"))
NTOP = int(os.environ.get("DG_NTOP", "30"))
ONLY = [b.strip().lower() for b in os.environ.get("DG_ONLY", "").split(",") if b.strip()]

COLS = tuple(x.strip() for x in os.environ.get(
    "DG_COLS",
    "productIdentificationName,productIdentificationTrademark,manufacterName,applicantName"
).split(",") if x.strip())
RDS_URL = "https://pub.fsa.gov.ru/rds/declaration/view/{id}/product"
ALNUM = re.compile(r"[^0-9a-zа-яё]")
ITEM_SPLIT = re.compile(r"[,;\s]+")
STOP = {"motors", "motor", "auto", "autos", "parts", "part", "group", "brand",
        "ltd", "inc", "spa", "gmbh", "corp", "the", "and", "co", "llc", "sa"}


def load_tnved():
    id2code = {}
    if os.path.exists(TNVED_TREE):
        tree = json.load(open(TNVED_TREE, encoding="utf-8"))
        id2code = {int(k): v.get("code") for k, v in tree.items()
                   if v.get("code") and str(v.get("code")).isdigit()}
    return id2code


def union_tnved(detail, id2code):
    idents = (detail.get("product") or {}).get("identifications") or []
    out = []
    for ident in idents:
        ids = ident.get("idTnveds") or []
        ids = [t.get("id") if isinstance(t, dict) else t for t in ids]
        for t in ids:
            if t in id2code:
                out.append(id2code[t])
    return ";".join(dict.fromkeys(out))


def search(col, q, cap):
    out = []
    page = 0
    while len(out) < cap:
        st = data = None
        for attempt in range(3):
            try:
                st, data = fsa_api.search_declarations(
                    [{"name": col, "search": q, "type": 0}], size=100, page=page)
            except Exception:
                st, data = None, None
            if st == 200 and isinstance(data, dict):
                break
            time.sleep(2.0 * (attempt + 1))
        if st != 200 or not isinstance(data, dict):
            break
        items = data.get("items") or []
        out += items
        if len(items) < 100:
            break
        page += 1
        time.sleep(0.4)
    return out


def brand_keys(b):
    keys = [b]
    for p in re.split(r"[/,&]+", b):
        p = p.strip()
        if p and p != b:
            keys.append(p)
    if os.environ.get("DG_EXACT") == "1":
        return keys
    variants = set()
    for k in list(keys):
        variants.add(re.sub(r"[^0-9a-zа-яё]", "", k))
        for w in re.split(r"[\s\-\u2013]+", k):
            if len(w) >= 4 and w not in STOP:
                variants.add(w)
    for v in variants:
        if v and v not in keys:
            keys.append(v)
    return keys


def brand_rx(b):
    return re.compile(r"(?<![0-9a-z])" + re.escape(b) + r"(?![0-9a-z])")


def main():
    wb = openpyxl.load_workbook(IN_FILE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    ci, bi, ni = header.index("Code"), header.index("BrandName"), header.index("CertNumber")
    rows = []
    empty = defaultdict(list)
    codes = defaultdict(set)
    for r in it:
        row = list(r)
        rows.append(row)
        b = str(row[bi] or "").strip().lower()
        if not b or b == "zekkert":
            continue
        if not (row[ni] and str(row[ni]).strip()):
            empty[b].append(len(rows) - 1)
            c = str(row[ci] or "").strip().lower()
            if c:
                codes[b].add(ALNUM.sub("", c))
    wb.close()
    brands = sorted(empty)
    if ONLY:
        brands = [b for b in brands if b in ONLY]
    print(f"input rows={len(rows)} empty brands={len(brands)} "
          f"empty rows={sum(len(empty[b]) for b in brands)}", flush=True)

    cache = json.load(open(CACHE, encoding="utf-8")) if os.path.exists(CACHE) else {}
    meta = {}

    for i, b in enumerate(brands, 1):
        if b in cache:
            continue
        s = set()
        for key in brand_keys(b):
            for col in COLS:
                try:
                    items = search(col, key, NLIVE)
                except Exception:
                    items = []
                for it2 in items:
                    did = str(it2.get("id"))
                    s.add(did)
                    meta.setdefault(did, {"number": it2.get("declNumber") or it2.get("number"),
                                          "date": it2.get("declDate"),
                                          "idStatus": it2.get("idStatus"),
                                          "applicantName": it2.get("applicantName")})
                time.sleep(0.3)
        cache[b] = sorted(s)
        json.dump(cache, open(CACHE, "w"), ensure_ascii=False)
        print(f"[{i}/{len(brands)}] {b}: cands={len(cache[b])}", flush=True)

    def cand_sort(did):
        m = meta.get(did, {})
        return (1 if m.get("idStatus") == 6 else 0, str(m.get("date") or ""))

    pool = {}
    need = set()
    for b in brands:
        cs = sorted(cache.get(b, []), key=cand_sort, reverse=True)[:NTOP]
        pool[b] = cs
        need.update(cs)

    cards = json.load(open(CARDS, encoding="utf-8")) if os.path.exists(CARDS) else {}
    missing = sorted(need - set(cards), key=lambda x: int(x))
    print(f"details needed={len(need)} missing={len(missing)}", flush=True)
    for j, did in enumerate(missing, 1):
        try:
            st, data = fsa_api.get_declaration(did)
        except Exception:
            continue
        if st == 200 and isinstance(data, dict):
            cards[did] = data
            if j % 25 == 0:
                json.dump(cards, open(CARDS, "w"), ensure_ascii=False)
        if j % 50 == 0:
            print(f"   details {j}/{len(missing)}", flush=True)
        time.sleep(0.2)
    json.dump(cards, open(CARDS, "w"), ensure_ascii=False)
    print("detail fetch done", flush=True)

    id2code = load_tnved()
    attach = {}
    for b in brands:
        rxs = [brand_rx(k) for k in brand_keys(b)]
        bcodes = codes.get(b, set())
        best, bestkey = None, None
        for did in pool.get(b, []):
            d = cards.get(did)
            if d is None:
                continue
            prod = d.get("product") or {}
            idents = prod.get("identifications") or []
            man = (d.get("manufacturer") or {}).get("fullName") or ""
            app = (d.get("applicant") or {}).get("fullName") or ""
            extra = B.norm_text(" ".join([str(prod.get("fullName") or ""), str(man), str(app)]))
            ntxt = (" ".join(B.norm_text(str(x.get("name") or "")) for x in idents)
                    + " " + extra).lower()
            ttxt = " ".join(B.norm_text(str(x.get("tradeMark") or "")) for x in idents).lower()
            atxt = " ".join(str(x.get("article") or "") for x in idents).lower()
            name_hit = any(r.search(ntxt) for r in rxs)
            tm_hit = any(r.search(ttxt) for r in rxs)
            items = {ALNUM.sub("", w) for w in ITEM_SPLIT.split(atxt) if w}
            art_hit = sum(1 for k in bcodes if len(k) >= 5 and k in items)
            key = (1 if art_hit else 0, 1 if name_hit else 0, 1 if tm_hit else 0,
                   str(d.get("declRegDate") or ""))
            if best is None or key > bestkey:
                best, bestkey = did, key
        if best is None or not (bestkey[0] or bestkey[1] or bestkey[2]):
            print(f"   {b}: no verified declaration (cands={len(pool.get(b,[]))})", flush=True)
            continue
        d = cards[best]
        number = d.get("number") or d.get("declNumber") or ""
        url = RDS_URL.format(id=best)
        tnved = B.cap_tnved(union_tnved(d, id2code))
        for ridx in empty[b]:
            attach[ridx] = (number, url, tnved, "Декларация", best)
        print(f"   {b:16} <- {best} art={bestkey[0]} name={bestkey[1]} tm={bestkey[2]} "
              f"rows={len(empty[b])} {str(number)[:34]}", flush=True)

    ni, ui, ti, yi = (header.index(x) for x in ("CertNumber", "CertUrl", "TNVED", "CertType"))
    st = Counter()
    for ridx, (number, url, tnved, ctype, did) in attach.items():
        rows[ridx][ni] = str(number)
        rows[ridx][ui] = url
        rows[ridx][ti] = tnved
        rows[ridx][yi] = ctype
        st["filled"] += 1
        if tnved:
            st["with_tnved"] += 1
    owb = openpyxl.Workbook(write_only=True)
    ows = owb.create_sheet("PRICE")
    ows.append(header)
    for r in rows:
        r = ["" if v is None else v for v in r]
        if len(r) < len(header):
            r += [""] * (len(header) - len(r))
        ows.append(r)
    owb.save(OUT_FILE)
    print(f"declaration fills={dict(st)}", flush=True)
    print(f"written {OUT_FILE} ({len(rows)} rows)", flush=True)


if __name__ == "__main__":
    main()
