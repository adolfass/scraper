#!/usr/bin/env python3
"""Diagnose TNVED recovery chain (Y ident -> FSA X ident -> survivor)
and cell-size distributions, with progress + slow-key logging."""
import json
import os
import re
import sys
import time
from collections import Counter

t0 = time.time()


def tick(msg):
    print(f"[{time.time() - t0:6.1f}s] {msg}", flush=True)


import build_output as bo

rc_bak = json.load(open("ref_carry_certs.json.bak", encoding="utf-8"))
rc_now = json.load(open("ref_carry_certs.json", encoding="utf-8"))
need = {k: v for k, v in rc_bak.items() if not v.get("tnved")}
tick(f"ref-empty entries: {len(need)}")
now_filled = sum(1 for k in need if rc_now.get(k, {}).get("tnved"))
tick(f"now filled among them: {now_filled}, "
     f"still empty: {len(need) - now_filled}")

log, rows, certs = bo.load_state()
tick(f"state loaded: probe={len(log)} rows={len(rows)} details={len(certs)}")
cand = bo.build_candidates(log)
chosen = {c: bo.choose(rids, rows) for c, rids in cand.items()}
tick(f"chosen computed: {len(chosen)} codes")
decls = json.load(open("ref_carry_decl_details.json", encoding="utf-8"))

# precompute per-detail indexes: token -> ident, and [(art, ident), ...]
IDX = {}
def detail_index(d):
    key = id(d)
    if key in IDX:
        return IDX[key]
    tok, arts = {}, []
    idents = (d.get("product") or {}).get("identifications") or []
    for ident in idents:
        art = (ident.get("article") or "").lower()
        if not art:
            continue
        arts.append((art, ident))
        for t in bo.pipeline.TOKEN_SPLIT.split(art):
            if t and t not in tok:
                tok[t] = ident
    IDX[key] = (tok, arts)
    return IDX[key]


def find_ident(d, code):
    tok, arts = detail_index(d)
    ident = tok.get(code)
    if ident is not None:
        return ident
    for art, ident in arts:
        if code in art:
            return ident
    return None


st = Counter()
slow = []
last = time.time()
for i, (k, v) in enumerate(need.items()):
    if i and i % 5000 == 0:
        tick(f"  {i}/{len(need)} {dict(st)} "
             f"(mem last {time.time() - last:.2f}s)")
    ks = time.time()
    u = v.get("url") or ""
    m = re.search(r"/view/([^/]+)", u)
    xid = m.group(1) if m else None
    code = k.split("|", 1)[1] if "|" in k else ""
    try:
        yd = (decls.get(xid) if "/rds/" in u else
              certs.get(int(xid)) if xid and xid.isdigit() else None)
        if yd and code and find_ident(yd, code):
            st["Y_ident"] += 1
        else:
            xc = chosen.get(code)
            xd = certs.get(xc) if xc is not None else None
            if xd and code and find_ident(xd, code):
                st["X_ident_fsa"] += 1
            else:
                st["survivor"] += 1
    except BaseException as e:
        st["ERROR"] += 1
        tick(f"  ERROR on {k}: {type(e).__name__}: {e}")
    dt = time.time() - ks
    if dt > 1.5:
        slow.append((round(dt, 2), k))
    last = time.time()
tick(f"recovery chain: {dict(st)}")
if slow:
    tick(f"slow keys (>1.5s): {len(slow)} top: {sorted(slow, reverse=True)[:5]}")
tick(f"distinct details indexed: {len(IDX)}")

import openpyxl


def col_dist(path, colname):
    ts = time.time()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    ci = hdr.index(colname)
    sizes = []
    for r in it:
        v = r[ci]
        if v:
            sizes.append(len(str(v).split(";")))
    wb.close()
    sizes.sort()
    res = None
    if sizes:
        def p(q):
            return sizes[min(len(sizes) - 1, int(q * len(sizes)))]
        res = {"n": len(sizes), "p50": p(.5), "p90": p(.9), "p99": p(.99),
               "max": sizes[-1], "gt10": sum(1 for s in sizes if s > 10),
               "gt30": sum(1 for s in sizes if s > 30)}
    tick(f"col_dist {path}: {time.time() - ts:.1f}s")
    return res


tick("REF  TNVED cells: " + str(col_dist("Готовый.xlsx", "TNVED")))
tick("OURS TNVED cells: " + str(col_dist("Результат.xlsx", "TNVED")))
tick("DONE")
sys.exit(0)
