#!/usr/bin/env python3
"""Find BRAND-LEVEL certificates: cards whose identification names a trade mark
and/or TU but does NOT enumerate specific article codes (the actual article list
lives in the annex).  Such a certificate can cover many of our rows for that
brand at once - the exact "one URL per several parts" case.

For every unmatched price-list row's brand we look for matching brand-level
cards and report the potential coverage.
"""
import json
import re
from collections import Counter, defaultdict

import openpyxl

import build_output as B

QUOTED = re.compile(r"[«\"']([^»\"']{2,60})[»\"']")
CODEISH = re.compile(r"\b[0-9a-z]{4,}\b")


def brand_text(d):
    parts = []
    prod = d.get("product") or {}
    parts += [str(prod.get("fullName") or ""), str(prod.get("name") or "")]
    for idt in prod.get("identifications") or []:
        parts += [str(idt.get("name") or ""), str(idt.get("description") or "")]
        tm = idt.get("tradeMark")
        if isinstance(tm, list):
            parts += [str(x) for x in tm]
        elif tm:
            parts.append(str(tm))
        for doc in idt.get("documents") or []:
            if isinstance(doc, dict):
                parts.append(str(doc.get("name") or ""))
    parts.append(str(d.get("tradeMark") or ""))
    return " ".join(parts)


def is_brand_level(d):
    """True when no identification enumerates a concrete article list."""
    prod = d.get("product") or {}
    articles = []
    for idt in prod.get("identifications") or []:
        a = str(idt.get("article") or "").strip()
        if a:
            articles.append(a)
    if not articles:
        return True
    # an article field enumerates codes if it has many long alnum tokens
    for a in articles:
        if len(CODEISH.findall(a.lower())) >= 3 and re.search(r"\d", a):
            return False
    return True


def main():
    # unmatched rows from shipped file
    wb = openpyxl.load_workbook("Итоговый ИП_Результат.xlsx",
                                read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    ci, bi, pi, ni = (h.index("Code"), h.index("BrandName"),
                      h.index("ProdName"), h.index("CertNumber"))
    un_rows = []
    for r in it:
        if not (r[ni] and str(r[ni]).strip()):
            un_rows.append((str(r[bi] or "").strip().lower(),
                            str(r[ci] or "").strip(),
                            str(r[pi] or "").strip()))
    wb.close()
    brands = sorted({b for b, _, _ in un_rows if b})
    rows_per_brand = Counter(b for b, _, _ in un_rows)
    print(f"unmatched rows={len(un_rows)} distinct brands={len(brands)}")

    per_brand_level = defaultdict(list)   # brand -> [(rid, num, prodtext)]
    per_brand_any = Counter()
    cards = 0
    for line in open("certs.jsonl", encoding="utf-8"):
        rec = json.loads(line)
        d = rec.get("data") or {}
        if not d:
            continue
        cards += 1
        bt = brand_text(d)
        nb = B.norm_text(bt)
        blevel = is_brand_level(d)
        prod = d.get("product") or {}
        pt = str(prod.get("fullName") or prod.get("name") or "")[:80]
        for b in brands:
            if B.brand_matches(b, nb):
                per_brand_any[b] += 1
                if blevel:
                    num = d.get("docId") or rec.get("id")
                    per_brand_level[b].append((rec.get("id"), num, pt))
    print(f"cards={cards}")

    print("\n== brands with BRAND-LEVEL certs (potential one-URL-many-parts) ==")
    tot_potential = 0
    rows = []
    for b in brands:
        bl = per_brand_level.get(b) or []
        if bl:
            rows.append((len(bl), per_brand_any[b], rows_per_brand[b], b))
            tot_potential += rows_per_brand[b]
    rows.sort(reverse=True)
    print(f"{'#bl':>4} {'#any':>5} {'#unrows':>7}  brand")
    for nbl, nany, nun, b in rows[:35]:
        print(f"{nbl:4d} {nany:5d} {nun:7d}  {b!r}")
    print(f"total unmatched rows whose brand has >=1 brand-level cert: "
          f"{tot_potential} / {len(un_rows)}")

    json.dump({b: v for b, v in per_brand_level.items()},
              open("/tmp/brand_level.json", "w"), ensure_ascii=False)


if __name__ == "__main__":
    main()