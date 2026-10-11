#!/usr/bin/env python3
"""Targeted scout: only download exports for cards whose text matches a brand
that still has UNMATCHED rows, then check how many of those unmatched codes
appear in the annex article lists."""
import io
import json
import re
import sys
import time
import zipfile
from collections import defaultdict

import openpyxl

import build_output as B
import fsa_api

CACHE = "/tmp/fsa_docx"
TOK = re.compile(r"[0-9a-zа-яё]+")


def docx_text(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", xml))


def get_export(cid):
    import os
    path = f"{CACHE}/{cid}.docx"
    if os.path.exists(path):
        return open(path, "rb").read()
    st, data = fsa_api._call(
        "POST", f"/api/v1/rss/common/certificates/{cid}/export/docx",
        body={}, retries=2)
    if st == 200 and isinstance(data, (bytes, bytearray)) and data[:2] == b"PK":
        os.makedirs(CACHE, exist_ok=True)
        open(path, "wb").write(data)
        return data
    return None


def card_text(d):
    parts = []
    prod = d.get("product") or {}
    parts += [str(prod.get("fullName") or ""), str(prod.get("name") or "")]
    for idt in prod.get("identifications") or []:
        parts += [str(idt.get("name") or ""), str(idt.get("description") or "")]
        tm = idt.get("tradeMark")
        parts += [str(x) for x in tm] if isinstance(tm, list) else ([str(tm)] if tm else [])
        for doc in idt.get("documents") or []:
            if isinstance(doc, dict):
                parts.append(str(doc.get("name") or ""))
    for who in ("applicant", "manufacturer"):
        o = d.get(who) or {}
        parts += [str(o.get("name") or ""), str(o.get("fullName") or "")]
    return " ".join(parts)


def main():
    limit = int(sys.argv[1]) if len(sys.argv) > 1 else 120
    wb = openpyxl.load_workbook("Итоговый ИП_Результат.xlsx",
                                read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    ci, bi, ni = h.index("Code"), h.index("BrandName"), h.index("CertNumber")
    by_brand = defaultdict(list)
    for r in it:
        code = str(r[ci] or "").strip().lower()
        if not code or (r[ni] and str(r[ni]).strip()):
            continue
        by_brand[str(r[bi] or "").strip().lower()].append(code)
    wb.close()
    brands = [b for b in by_brand if b]
    print(f"unmatched brands={len(brands)} codes={sum(len(v) for v in by_brand.values())}")

    # select cards matching a brand that has unmatched rows
    match_cards = []
    for line in open("certs.jsonl", encoding="utf-8"):
        rec = json.loads(line)
        d = rec.get("data") or {}
        if not d:
            continue
        nb = B.norm_text(card_text(d))
        for b in brands:
            if B.brand_matches(b, nb):
                match_cards.append((rec.get("id"), b, d))
                break
    print(f"cards matching an unmatched brand: {len(match_cards)}")
    # de-dup keep one brand per card
    seen = set()
    uniq = []
    for rid, b, d in match_cards:
        if rid in seen:
            continue
        seen.add(rid)
        uniq.append((rid, b, d))

    hits = defaultdict(list)
    done = 0
    t0 = time.time()
    for rid, b, d in uniq[:limit]:
        data = get_export(rid)
        done += 1
        if not data:
            continue
        toks = set(TOK.findall(docx_text(data).lower()))
        found = toks & set(by_brand.get(b, []))
        for c in found:
            hits[b].append((c, rid))
        if done % 10 == 0:
            print(f"  [{done}/{limit}] {time.time()-t0:.0f}s hits={sum(len(v) for v in hits.values())}",
                  flush=True)
        time.sleep(0.3)
    print(f"\ndone {done} in {time.time()-t0:.0f}s")
    tot = sum(len(v) for v in hits.values())
    codes = {c for v in hits.values() for c, _ in v}
    print(f"unmatched codes found in annex exports: {len(codes)} (pairs={tot})")
    for b, v in sorted(hits.items(), key=lambda x: -len(x[1]))[:25]:
        print(f"   {b:18} +{len(v):4}  e.g. {[c for c,_ in v[:6]]}")
    json.dump({b: v for b, v in hits.items()},
              open("/tmp/export_scout_targeted.json", "w"), ensure_ascii=False)


if __name__ == "__main__":
    main()
