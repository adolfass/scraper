#!/usr/bin/env python3
"""Merge brand_group shard outputs into one workbook and QA the result.

Env:
  MG_IN     baseline workbook (default /tmp/Итоговый_ИП_brand2.xlsx)
  MG_PARTS  comma list of shard outputs
  MG_OUT    merged output (default /tmp/Итоговый_ИП_final.xlsx)
"""
import os
import sys

import openpyxl

IN = os.environ.get("MG_IN", "/tmp/Итоговый_ИП_brand2.xlsx")
PARTS = [p for p in os.environ.get(
    "MG_PARTS",
    "/tmp/Итоговый_ИП_group_0.xlsx,/tmp/Итоговый_ИП_group_1.xlsx,"
    "/tmp/Итоговый_ИП_group_2.xlsx").split(",") if p]
OUT = os.environ.get("MG_OUT", "/tmp/Итоговый_ИП_final.xlsx")


def load(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    rows = [list(r) for r in it]
    wb.close()
    return h, rows


def main():
    h, base = load(IN)
    ni, ui, ti, yi = (h.index(x) for x in ("CertNumber", "CertUrl", "TNVED", "CertType"))
    base_filled = sum(1 for r in base if r[ni] and str(r[ni]).strip())
    changed = 0
    added = 0
    for p in PARTS:
        if not os.path.exists(p):
            print("missing shard:", p, file=sys.stderr)
            continue
        ph, prows = load(p)
        assert ph == h, f"header mismatch in {p}"
        assert len(prows) == len(base), f"row count mismatch in {p}"
        for i in range(len(base)):
            had = base[i][ni] and str(base[i][ni]).strip()
            new = prows[i][ni] and str(prows[i][ni]).strip()
            if not had and new:
                base[i][ni], base[i][ui] = prows[i][ni], prows[i][ui]
                base[i][ti], base[i][yi] = prows[i][ti], prows[i][yi]
                added += 1
            elif had and new and str(base[i][ni]) != str(prows[i][ni]):
                changed += 1
    owb = openpyxl.Workbook(write_only=True)
    ows = owb.create_sheet("PRICE")
    ows.append(h)
    zek = 0
    filled = 0
    for r in base:
        r = ["" if v is None else v for v in r]
        if len(r) < len(h):
            r += [""] * (len(h) - len(r))
        if r[ni] and str(r[ni]).strip():
            filled += 1
        if str(r[h.index("BrandName")] or "").strip().lower() == "zekkert":
            zek += 1
        ows.append(r)
    owb.save(OUT)
    print(f"baseline filled={base_filled}  shard additions={added}  "
          f"conflicts(ignored)={changed}")
    print(f"final: rows={len(base)} filled={filled} ({100*filled/len(base):.1f}%) "
          f"zekkert={zek}")
    print(f"written {OUT}")


if __name__ == "__main__":
    main()
