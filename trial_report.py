#!/usr/bin/env python3
"""Trial report: first N source rows of Результат.xlsx vs Готовый.xlsx.

Same metrics as build_output.py validate, but restricted to the first N
rows of the price list in file order (default 500).

  .venv/bin/python trial_report.py [N]
"""
import sys
import openpyxl

ROOT = __file__.rsplit("/", 1)[0]
REF = f"{ROOT}/Готовый.xlsx"
OUT = f"{ROOT}/Результат.xlsx"
SRC = f"{ROOT}/Исходный.xlsx"
COLS = ("CertNumber", "CertUrl", "TNVED", "CertType")


def load(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    idx = {h: i for i, h in enumerate(hdr)}
    res = {}
    for r in it:
        res[(str(r[idx["ProdNum"]]), str(r[idx["Code"]]))] = {
            c: (str(r[idx[c]]).strip() if r[idx[c]] else "") for c in COLS}
    wb.close()
    return res


def first_keys(n):
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["PRICE"]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    pcol = "ProdNum" if "ProdNum" in hdr else "Prodnum"
    idx = {h: i for i, h in enumerate(hdr)}
    ks = []
    for r in it:
        ks.append((str(r[idx[pcol]]), str(r[idx["Code"]])))
        if len(ks) >= n:
            break
    wb.close()
    return ks


def main():
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 500
    keys = first_keys(n)
    print(f"trial: first {n} rows of the price list", flush=True)
    ref, ours = load(REF), load(OUT)

    from collections import Counter
    st = Counter()
    agreed, differ = [], []
    for k in keys:
        rc = ref.get(k, {}).get("CertNumber", "")
        oc = ours.get(k, {}).get("CertNumber", "")
        if rc and oc:
            if rc == oc:
                st["cert:agree"] += 1
                agreed.append(k)
            else:
                st["cert:differ"] += 1
                differ.append(k)
        elif rc:
            st["cert:ref_only"] += 1
        elif oc:
            st["cert:ours_only"] += 1
        else:
            st["cert:both_empty"] += 1
    for k in keys:
        if k not in ref:
            st["row:not_in_ref"] += 1

    print("\nCertNumber:")
    for k, v in sorted(st.items()):
        print(f"  {k:22s} {v}")

    t = Counter()
    for k in agreed:
        r, o = ref[k], ours[k]
        rt, ot = r["TNVED"], o["TNVED"]
        if rt and ot:
            t["tnved:agree" if rt == ot else "tnved:differ"] += 1
        elif rt:
            t["tnved:ref_only"] += 1
        elif ot:
            t["tnved:ours_only"] += 1
        else:
            t["tnved:both_empty"] += 1
        if r["CertType"] and o["CertType"]:
            t["type:agree" if r["CertType"] == o["CertType"] else "type:differ"] += 1
    print("\non agreed-cert rows:")
    for k, v in sorted(t.items()):
        print(f"  {k:22s} {v}")

    if differ:
        print(f"\ncert mismatch samples ({len(differ)}):")
        for k in differ[:10]:
            print(f"   {k[1]:16s} ref: {ref[k]['CertNumber']} | ours: {ours[k]['CertNumber']}")
    if st["cert:ref_only"]:
        ro = [k for k in keys
              if ref.get(k, {}).get("CertNumber") and not ours.get(k, {}).get("CertNumber")]
        print(f"\nref-only samples ({len(ro)}):")
        for k in ro[:10]:
            print(f"   {k[1]:16s} ref: {ref[k]['CertNumber']} | {ref[k]['CertUrl'][:70]}")


if __name__ == "__main__":
    main()
