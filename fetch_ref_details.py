#!/usr/bin/env python3
"""Fetch certificate details for every FSA cert referenced in Готовый.xlsx.

Gentle + resumable -> ref_certs.jsonl ({id, data}).
Serves two purposes:
  1) distribution of idTechnicalReglaments across reference certs
     (to evaluate a registry scan filtered by technical regulation);
  2) details cached for later use (TNVED etc.).
"""
import json, os, re, time, sys
import openpyxl, fsa_api

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "ref_certs.jsonl")
DELAY = 0.7


def ref_fsa_ids():
    wb = openpyxl.load_workbook(os.path.join(ROOT, "Готовый.xlsx"),
                                read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    idx = {x: i for i, x in enumerate(h)}
    ids = set()
    for r in it:
        u = r[idx["CertUrl"]] or ""
        m = re.search(r"/view/(\d+)/", str(u))
        if m:
            ids.add(int(m.group(1)))
    wb.close()
    return sorted(ids)


def have():
    out = set()
    if os.path.exists(OUT):
        with open(OUT, encoding="utf-8") as f:
            for line in f:
                try:
                    out.add(json.loads(line)["id"])
                except Exception:
                    pass
    return out


def main():
    ids = ref_fsa_ids()
    done = have()
    todo = [i for i in ids if i not in done]
    print(f"ref FSA cert ids: {len(ids)}, cached: {len(done)}, todo: {len(todo)}",
          flush=True)
    for n, cid in enumerate(todo, 1):
        for attempt in range(4):
            try:
                st, data = fsa_api.get_certificate(cid)
                if st == 200:
                    break
                raise RuntimeError(f"http {st}")
            except Exception as e:
                wait = min(5 * 2 ** attempt, 120)
                print(f"  cert {cid} attempt {attempt+1} failed: {e}; "
                      f"sleep {wait}s", flush=True)
                time.sleep(wait)
        else:
            print(f"  cert {cid} GAVE UP", flush=True)
            continue
        with open(OUT, "a", encoding="utf-8") as f:
            f.write(json.dumps({"id": cid, "data": data}, ensure_ascii=False) + "\n")
        time.sleep(DELAY)
        if n % 25 == 0:
            print(f"  {n}/{len(todo)}", flush=True)
    print("done", flush=True)


if __name__ == "__main__":
    main()
