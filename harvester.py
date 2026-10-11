#!/usr/bin/env python3
"""Harvest annex article lists by TRADE MARK.

For a brand, search all certificates whose productIdentificationTrademark
matches, download each certificate's print-form export (which contains the annex
article list), and check whether our still-unmatched codes for that brand appear.

Usage: pt harvester.py BRAND [BRAND ...]
"""
import io
import json
import os
import re
import sys
import time
import zipfile

import openpyxl

import fsa_api

CACHE = "/tmp/fsa_docx"
TOK = re.compile(r"[0-9a-zа-яё]+")
TRADEMARK_COL = "productIdentificationTrademark"


def docx_text(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", xml))


def get_export(cid):
    path = f"{CACHE}/{cid}.docx"
    if os.path.exists(path):
        return open(path, "rb").read()
    st, data = fsa_api._call(
        "POST", f"/api/v1/rss/common/certificates/{cid}/export/docx",
        body={}, retries=2)
    if st == 200 and isinstance(data, (bytes, bytearray)) and data[:2] == b"PK":
        os.makedirs(CACHE, exist_ok=True)
        open(path, "wb").write(data)
        return data
    return None


def search_trademark(brand, cap=200):
    ids, page, total = [], 0, None
    while True:
        st, data = fsa_api._call("POST", "/api/v1/rss/common/certificates/get",
                                 body={"size": 100, "page": page,
                                       "filter": {"idCertScheme": [],
                                                  "regDate": {"startDate": None, "endDate": None},
                                                  "endDate": {"startDate": None, "endDate": None},
                                                  "columnsSearch": [{"column": TRADEMARK_COL, "search": brand, "type": 0}]},
                                       "columnsSort": [{"column": "date", "sort": "DESC"}]},
                                 retries=2)
        if st != 200 or not isinstance(data, dict):
            break
        items = data.get("items") or []
        if total is None:
            total = data.get("totalCount") or data.get("total") or len(items)
        for it in items:
            ids.append(it.get("id"))
        if len(items) < 100 or len(ids) >= cap:
            break
        page += 1
        time.sleep(0.3)
    return total, ids


def main():
    want = [a.upper() for a in sys.argv[1:]] or ["TOPRAN", "MAHLE", "NGK", "DENSO",
                                                 "GATES", "SKF", "STELLOX", "FILTRON",
                                                 "RUVILLE", "FEBI", "MEYLE", "LEMFORDER",
                                                 "FEBEST", "TIMKEN", "ELRING", "SWAG",
                                                 "PATRON", "VICTOR REINZ", "PIERBURG",
                                                 "JP GROUP", "AUTOMEGA", "VEMO"]
    wb = openpyxl.load_workbook("Итоговый ИП_Результат.xlsx",
                                read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    ci, bi, ni = h.index("Code"), h.index("BrandName"), h.index("CertNumber")
    by_brand = {}
    for r in it:
        code = str(r[ci] or "").strip().lower()
        if not code or (r[ni] and str(r[ni]).strip()):
            continue
        by_brand.setdefault(str(r[bi] or "").strip().lower(), set()).add(code)
    wb.close()

    all_hits = {}
    for brand in want:
        key = brand.lower()
        codes = by_brand.get(key)
        if not codes:
            # try loose: brand as substring of a price-list brand
            cand = [b for b in by_brand if key in b]
            if len(cand) == 1:
                key = cand[0]
                codes = by_brand[key]
        if not codes:
            print(f"{brand:14} no unmatched price rows")
            continue
        t0 = time.time()
        total, ids = search_trademark(brand)
        print(f"{brand:14} unmatched={len(codes)} certs_total={total} ids={len(ids)} ({time.time()-t0:.0f}s)")
        hits = 0
        for cid in ids:
            data = get_export(cid)
            if not data:
                continue
            toks = set(TOK.findall(docx_text(data).lower()))
            found = toks & codes
            if found:
                hits += len(found)
                for c in found:
                    all_hits.setdefault(c, []).append(cid)
                print(f"      cert {cid}: +{len(found)} {sorted(found)[:8]}")
            time.sleep(0.2)
        print(f"   -> {brand}: filled candidates={hits}")
    json.dump(all_hits, open("/tmp/harvest_hits.json", "w"), ensure_ascii=False)
    print(f"\nTOTAL distinct unmatched codes found: {len(all_hits)}")


if __name__ == "__main__":
    main()
