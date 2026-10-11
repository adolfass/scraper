#!/usr/bin/env python3
"""Scout: download the print-form export (which contains the ANNEX article
lists) for a sample of cards and check whether UNMATCHED price-list codes appear
there.  This validates the "export unlocks the annex" idea end-to-end.

Cache: /tmp/fsa_docx/<certid>.docx
"""
import io
import json
import os
import re
import sys
import time
import zipfile

import openpyxl

import fsa_api

CACHE = "/tmp/fsa_docx"
TOK = re.compile(r"[0-9a-zа-яё]+")


def docx_text(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    txt = re.sub(r"<[^>]+>", " ", xml)
    return re.sub(r"\s+", " ", txt)


def get_export(cid):
    path = f"{CACHE}/{cid}.docx"
    if os.path.exists(path):
        return open(path, "rb").read()
    st, data = fsa_api._call(
        "POST", f"/api/v1/rss/common/certificates/{cid}/export/docx",
        body={}, retries=2)
    if st == 200 and isinstance(data, (bytes, bytearray)) and data[:2] == b"PK":
        os.makedirs(CACHE, exist_ok=True)
        open(path, "wb").write(data)
        return data
    return None


def main():
    limit = int(sys.argv[1]) if len(sys.argv) > 1 else 60
    wb = openpyxl.load_workbook("Итоговый ИП_Результат.xlsx",
                                read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    ci, bi, ni = h.index("Code"), h.index("BrandName"), h.index("CertNumber")
    unmatched = {}
    for r in it:
        code = str(r[ci] or "").strip().lower()
        if not code or (r[ni] and str(r[ni]).strip()):
            continue
        unmatched.setdefault(code, str(r[bi] or "").strip().lower())
    wb.close()
    print(f"unmatched codes={len(unmatched)}")

    # pick annex-only cards
    cards = []
    for line in open("certs.jsonl", encoding="utf-8"):
        rec = json.loads(line)
        d = rec.get("data") or {}
        if not d:
            continue
        prod = d.get("product") or {}
        arts = [str(i.get("article") or "") for i in prod.get("identifications") or []]
        if any("приложени" in a.lower() for a in arts):
            cards.append((rec.get("id"), d))
    print(f"annex-only cards available={len(cards)}; sampling {limit}")

    hits = {}
    done = 0
    t0 = time.time()
    for rid, d in cards[:limit]:
        data = get_export(rid)
        done += 1
        if not data:
            print(f"  [{done}] cert {rid}: export failed")
            continue
        text = docx_text(data)
        toks = set(TOK.findall(text.lower()))
        found = toks & set(unmatched)
        for c in found:
            hits.setdefault(c, []).append(rid)
        if found:
            print(f"  [{done}] cert {rid}: +{len(found)} unmatched "
                  f"e.g. {sorted(found)[:6]}")
        time.sleep(0.4)
    print(f"\ndone {done} cards in {time.time()-t0:.0f}s")
    print(f"unmatched codes found in exports: {len(hits)} / {len(unmatched)}")
    json.dump(hits, open("/tmp/export_scout_hits.json", "w"), ensure_ascii=False)
    for c, rids in list(hits.items())[:25]:
        print(f"   {c:16} brand={unmatched[c]:14} certs={rids[:3]}")


if __name__ == "__main__":
    main()
