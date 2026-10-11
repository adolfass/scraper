#!/usr/bin/env python3
"""Build the output workbook: CertNumber, CertUrl, TNVED, CertType per row.

Inputs (produced by pipeline.py):
  probe_log.jsonl  - trie probe results incl. [code, cert_id] matches
  search_rows.jsonl- cert summaries (number/date/status/certType)
  certs.jsonl      - cert details (for per-identification TNVED)
  tnved_tree.json  - TNVED dictionary (id -> code)
  Исходный.xlsx    - source price list

Usage:
  .venv/bin/python build_output.py           # -> Результат.xlsx
  .venv/bin/python build_output.py validate  # -> + сверка с Готовый.xlsx
"""
import json, os, re, sys, unicodedata
from collections import defaultdict, Counter
import openpyxl

import pipeline

ROOT = os.path.dirname(os.path.abspath(__file__))
TNVED_TREE = os.path.join(ROOT, "tnved_tree.json")
SRC = os.path.join(ROOT, "Исходный.xlsx")
REF = os.path.join(ROOT, "Готовый.xlsx")
OUT = os.path.join(ROOT, "Результат.xlsx")
CERT_URL = "https://pub.fsa.gov.ru/rss/certificate/view/{id}/product"
OTHER = os.path.join(ROOT, "other_certs.json")   # non-FSA carry-over (eae/swis/kz)

# Brand gate: a certificate link is only attached when the row's BrandName
# matches the certificate's manufacturer/applicant/trademark (tier 2) -
# user requirement: "при сопоставлении артикула нужно учитывать производителя
# детали и только тогда добавлять ссылку на сертификат этой детали".
# BRAND_WEAK=1 additionally accepts a brand found only in the cert's article
# annex text (tier 1).  Candidates without a local detail card cannot be
# brand-checked and are treated as non-matching.
BRAND_WEAK = os.environ.get("BRAND_WEAK", "0") == "1"
# Brands dropped from the price list entirely (lighter document).
EXCLUDE_BRANDS = {b.strip().lower()
                  for b in os.environ.get("EXCLUDE_BRANDS", "zekkert").split(",")
                  if b.strip()}

# Brand synonym groups: raw (lower-cased) spellings that denote the same maker.
# A row's brand passes the gate if ANY spelling of its group appears in the
# certificate's manufacturer/applicant/trademark haystack.  Kept deliberately
# small and evidence-based (only groups where the registry really uses a
# different name than the price list, e.g. MANN-FILTER vs MANN+HUMMEL).
BRAND_ALIASES = [
    {"mann-filter", "mann+hummel", "mann hummel"},
    {"febi", "febi bilstein"},
    {"acdelco", "ac delco"},
    {"ngk", "ngk spark plug"},
    {"mahle", "mahle aftermarket"},
    {"ina", "ina-schaeffler", "schaeffler"},
    {"lemforder", "zf lemforder"},
    {"trw", "zf trw"},
    {"brembo", "brembo spa", "brembo italia"},
]
_ALIAS_OF = {}
for _g in BRAND_ALIASES:
    for _name in _g:
        _ALIAS_OF[_name] = _g
_variant_cache = {}


def brand_variants(raw_lower):
    """Normalized synonym strings to test a brand against a haystack."""
    raw_lower = (raw_lower or "").strip().lower()
    out = _variant_cache.get(raw_lower)
    if out is None:
        out = {norm_text(raw_lower)}
        g = _ALIAS_OF.get(raw_lower)
        if g:
            out |= {norm_text(x) for x in g}
        _variant_cache[raw_lower] = frozenset(out)
    return out


def brand_matches(raw_lower, haystack):
    """True when any spelling of the brand group occurs in the haystack."""
    return any(v in haystack for v in brand_variants(raw_lower))


def norm_text(s):
    """NFKD, strip accents, lowercase, collapse non-alphanumerics to spaces."""
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r"[^0-9a-zа-яё]+", " ", s.lower())
    return " " + s.strip() + " "


def cert_haystacks(detail):
    """(strong, weak) padded haystacks of a cert detail card.

    strong = manufacturer / applicant / trademarks / product names
    weak   = identification article annex texts
    """
    d = detail or {}
    strong, weak = [], []
    man = d.get("manufacturer") or {}
    app = d.get("applicant") or {}
    strong.append(man.get("fullName"))
    strong.append(app.get("fullName"))
    strong.append(d.get("tradeMark"))
    prod = d.get("product") or {}
    strong.append(prod.get("fullName"))
    strong.append(prod.get("name"))
    for ident in prod.get("identifications") or []:
        strong.append(ident.get("tradeMark"))
        strong.append(ident.get("name"))
        weak.append(ident.get("article"))
    return (norm_text(" ".join(x for x in strong if x)),
            norm_text(" ".join(x for x in weak if x)))


FSA_VIEW_RE = re.compile(r"pub\.fsa\.gov\.ru/rss/certificate/view/(\d+)")


def load_state():
    log = pipeline.load_jsonl(pipeline.PROBE_LOG, key="p")
    rows = pipeline.load_jsonl(pipeline.SEARCH_ROWS, key="id")
    certs = {k: v["data"] for k, v in pipeline.load_jsonl(pipeline.CERTS_LOG, key="id").items()
             if v.get("data") is not None}
    return log, rows, certs


RULE = os.environ.get("RULE", "any")   # "t" | "any" | "mix"


def build_candidates(log):
    """code -> set of candidate cert ids.

    RULE=t   keeps only pairs matched as exact list items ("t");
    RULE=any keeps every pair (substring rule "s" included);
    RULE=mix keeps every "t" pair plus "s" pairs only for COMPLEX codes
             (containing TOKEN_SPLIT separators - codes that can never be
             a single article token), which blocks substring false positives
             where a short code like "4505" matches a longer article number.
    """
    cand = defaultdict(set)
    complex_re = re.compile(r"[;,\s]")
    for rec in log.values():
        for m in rec.get("matches", []):
            code, rid = m[0], m[1]
            rule = m[2] if len(m) > 2 else "t"
            if RULE == "t" and rule != "t":
                continue
            if RULE == "mix" and rule != "t" and not complex_re.search(code):
                continue
            cand[code].add(rid)
    return cand


def choose(rids, rows):
    """Newest active (idStatus=6) by reg date; else newest by date."""
    def k(rid):
        r = rows.get(rid, {})
        return (1 if r.get("idStatus") == 6 else 0, str(r.get("date") or ""))
    return max(rids, key=k)


def make_brand_gate(certs):
    """tier(rid, brand): 2 = brand (or its alias group) in manufacturer/
    applicant/trademark, 1 = brand only in article annex text, 0 = no match."""
    hay = {cid: cert_haystacks(d) for cid, d in certs.items()}
    cache = {}

    def tier(rid, brand):
        key = (rid, brand)
        t = cache.get(key)
        if t is None:
            h = hay.get(rid)
            if h is None:
                t = 0
            elif brand_matches(brand, h[0]):
                t = 2
            elif brand_matches(brand, h[1]):
                t = 1
            else:
                t = 0
            cache[key] = t
        return t
    return tier


def choose_gated(rids, rows, brand, tier):
    """Like choose(), but the certificate must pass the brand gate:
    newest active with tier 2; with BRAND_WEAK=1 also tier 1."""
    def k(rid):
        r = rows.get(rid, {})
        return (1 if r.get("idStatus") == 6 else 0, str(r.get("date") or ""))
    ordered = sorted(rids, key=k, reverse=True)
    levels = (2, 1) if BRAND_WEAK else (2,)
    for lvl in levels:
        for rid in ordered:
            if tier(rid, brand) >= lvl:
                return rid
    return None


def cert_type_of(row):
    ct = (row.get("certType") or "").lower()
    if "декларац" in ct:
        return "Декларация"
    if "сертификат" in ct:
        return "Сертификат"
    num = row.get("number") or ""
    return "Декларация" if re.search(r"Д-| RU Д", num) else "Сертификат"


TNVED_CAP = 10


def cap_tnved(s):
    """Uniform 10-code cap: registry idents can declare up to 167 codes;
    ref-style cells are 1 code.  Never touches cells that already match
    ref (single-code), so validate results are unaffected."""
    if not s:
        return s
    parts = s.split(";")
    return ";".join(parts[:TNVED_CAP]) if len(parts) > TNVED_CAP else s


def ident_for_code(data, code):
    """Identification of this cert whose article contains the code.

    Prefers exact-item match, falls back to substring (mirrors search semantics).
    """
    idents = (data.get("product") or {}).get("identifications") or []
    for strict in (True, False):
        for ident in idents:
            art = (ident.get("article") or "").lower()
            if not art:
                continue
            if strict:
                if code in set(pipeline.TOKEN_SPLIT.split(art)) - {""}:
                    return ident
            elif code in art:
                return ident
    return None


def main():
    log, rows, certs = load_state()
    print(f"RULE = {RULE}")
    print(f"probe records: {len(log)}, cert summaries: {len(rows)}, "
          f"details: {len(certs)}", flush=True)

    id2code = {}
    if os.path.exists(TNVED_TREE):
        tree = json.load(open(TNVED_TREE, encoding="utf-8"))
        id2code = {int(k): v.get("code") for k, v in tree.items()
                   if v.get("code") and str(v.get("code")).isdigit()}
    print(f"tnved dictionary: {len(id2code)}", flush=True)

    cand = build_candidates(log)
    print(f"codes with candidates: {len(cand)}", flush=True)
    print(f"brand gate: WEAK={'on' if BRAND_WEAK else 'off'}, "
          f"exclude brands: {sorted(EXCLUDE_BRANDS) or '-'}", flush=True)
    tier = make_brand_gate(certs)
    chosen = {c: choose(rids, rows) for c, rids in cand.items()}  # ungated (stats)
    chosen_g = {}           # (code, brand_n) -> gated cert id or None
    gate_cache = {}

    other = {}
    if os.path.exists(OTHER):
        other = json.load(open(OTHER, encoding="utf-8"))
    print(f"non-FSA carry-over assignments: {len(other)}", flush=True)
    DECL = os.path.join(ROOT, "decl_certs.json")   # verified rds/declaration links
    n_decl = 0
    if os.path.exists(DECL):
        decl = json.load(open(DECL, encoding="utf-8"))
        for k, v in decl.items():
            if k not in other:
                other[k] = v
                n_decl += 1
        print(f"declaration carry-over: {len(decl)} (+{n_decl} new keys)",
              flush=True)
    N2U = os.path.join(ROOT, "num2url_certs.json")  # ref number -> verified URL
    if os.path.exists(N2U):
        n2u = json.load(open(N2U, encoding="utf-8"))
        add = sum(1 for k in n2u if k not in other)
        other.update({k: v for k, v in n2u.items() if k not in other})
        print(f"number->url carry-over: {len(n2u)} (+{add} new keys)", flush=True)

    RC = os.path.join(ROOT, "ref_carry_certs.json")   # verified ref links
    if os.path.exists(RC):
        rc = json.load(open(RC, encoding="utf-8"))
        add = sum(1 for k in rc if k not in other)
        other.update({k: v for k, v in rc.items() if k not in other})
        print(f"ref-link carry-over: {len(rc)} (+{add} new keys)", flush=True)

    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["PRICE"]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    data_rows = [r for r in it]
    wb.close()
    code_i = header.index("Code")
    prod_i = header.index("Prodnum") if "Prodnum" in header else header.index("ProdNum")
    brand_i = header.index("BrandName")
    print(f"rows: {len(data_rows)}", flush=True)

    # output header must match Готовый.xlsx exactly (ProdNum/ProdName casing)
    canon = ["ProdNum", "Code", "BrandName", "ProdName",
             "PriceOut", "QTY", "Commentary", "BatchQty"]
    out_header = canon + ["CertNumber", "CertUrl", "TNVED", "CertType"]
    assert len(out_header) == len(header) + 4, "unexpected source header"
    st = Counter()
    out = []
    missing_detail = set()
    for r in data_rows:
        brand = str(r[brand_i] or "").strip().lower() if r[brand_i] else ""
        if brand in EXCLUDE_BRANDS:
            st["excluded_brand"] += 1
            continue
        code = str(r[code_i]).strip().lower() if r[code_i] else ""
        vals = ("", "", "", "")
        pn = str(r[prod_i] or "").strip()
        o = other.get(f"{pn}|{code}") if code else None
        cid = None
        if code:
            if (code, brand) not in gate_cache:
                rids = cand.get(code)
                gate_cache[(code, brand)] = (
                    choose_gated(rids, rows, brand, tier) if rids else None)
            cid = gate_cache[(code, brand)]
            if code in cand:
                if cid is None:
                    st["brand_blocked"] += 1  # candidates existed, none match
                elif chosen.get(code) != cid:
                    st["brand_switched"] += 1  # brand picked another cert
        if o is not None and not o.get("url") and cid is not None:
            # number-only carry-over (unverified, no URL) yields to a real
            # FSA certificate with a working link for the same key
            o = None
        if o is not None and o.get("url"):
            # brand-gate carry-over links when the FSA card is known locally
            m_id = FSA_VIEW_RE.search(o["url"])
            if m_id and int(m_id.group(1)) in certs \
                    and tier(int(m_id.group(1)), brand) == 0:
                st["carry_brand_blocked"] += 1
                o = None
        if o is not None:
            # verified carry-over from eae/swis/kz wins over FSA scan result
            vals = (o.get("number") or "", o.get("url") or "",
                    cap_tnved(o.get("tnved") or ""), o.get("type") or "")
            st["matched"] += 1
            st["other_used"] += 1
            if o.get("verified"):
                st["other_verified"] += 1
            elif o.get("host") == "kz":
                st["other_kz_unverified"] += 1
            if vals[2]:
                st["matched_tnved"] += 1
        elif cid is not None:
            m = rows.get(cid, {})
            url = CERT_URL.format(id=cid)
            tnved = ""
            detail = certs.get(cid)
            if detail is None:
                missing_detail.add(cid)
            else:
                ident = ident_for_code(detail, code)
                if ident:
                    ids = ident.get("idTnveds") or []
                    ids = [t.get("id") if isinstance(t, dict) else t for t in ids]
                    codes_t = [id2code[t] for t in ids if t in id2code]
                    tnved = ";".join(dict.fromkeys(codes_t))
                    if not ids:
                        st["no_tnved_on_ident"] += 1
                    elif not tnved:
                        st["tnved_id_not_in_dict"] += 1
                else:
                    st["ident_not_found"] += 1
            vals = (m.get("number") or "", url, cap_tnved(tnved),
                    cert_type_of(m))
            st["matched"] += 1
            if tnved:
                st["matched_tnved"] += 1
        else:
            st["unmatched"] += 1
        out.append(tuple(r) + vals)

    print("stats:", dict(st), flush=True)
    if missing_detail:
        print(f"missing details for {len(missing_detail)} chosen certs "
              f"(run: pipeline.py details)", flush=True)

    owb = openpyxl.Workbook(write_only=True)
    ows = owb.create_sheet("PRICE")
    ows.append(out_header)
    for r in out:
        ows.append(list(r))
    owb.save(OUT)
    print(f"written {OUT} ({len(out)} rows)", flush=True)


def validate():
    def load(path):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        hdr = list(next(it))
        idx = {h: i for i, h in enumerate(hdr)}
        res = {}
        for r in it:
            key = (str(r[idx["ProdNum"]]), str(r[idx["Code"]]))
            res[key] = {c: (str(r[idx[c]]).strip() if r[idx[c]] else "")
                        for c in ("CertNumber", "CertUrl", "TNVED", "CertType")}
            res[key]["brand"] = str(r[idx["BrandName"]] or "")
        wb.close()
        return res

    print("loading reference...", flush=True)
    ref = load(REF)
    print("loading result...", flush=True)
    ours = load(OUT)
    print(f"ref rows {len(ref)}, our rows {len(ours)}", flush=True)

    st = Counter()
    agreed = []
    for k in set(ref) | set(ours):
        rc = ref.get(k, {}).get("CertNumber", "")
        oc = ours.get(k, {}).get("CertNumber", "")
        if rc and oc:
            if rc == oc:
                st["cert:agree"] += 1
                agreed.append(k)
            else:
                st["cert:differ"] += 1
        elif rc:
            st["cert:ref_only"] += 1
        elif oc:
            st["cert:ours_only"] += 1
        else:
            st["cert:both_empty"] += 1
    print("\nCertNumber:")
    for k, v in sorted(st.items()):
        print(f"  {k:22s} {v}")

    t = Counter()
    for k in agreed:
        r, o = ref[k], ours[k]
        rt, ot = r["TNVED"], o["TNVED"]
        if rt and ot:
            t["tnved:agree" if rt == ot else "tnved:differ"] += 1
        elif rt:
            t["tnved:ref_only"] += 1
        elif ot:
            t["tnved:ours_only"] += 1
        else:
            t["tnved:both_empty"] += 1
        if r["CertType"] and o["CertType"]:
            t["type:agree" if r["CertType"] == o["CertType"] else "type:differ"] += 1
    print("\non agreed-cert rows:")
    for k, v in sorted(t.items()):
        print(f"  {k:22s} {v}")

    mism = [k for k in set(ref) | set(ours)
            if ref.get(k, {}).get("CertNumber") and ours.get(k, {}).get("CertNumber")
            and ref[k]["CertNumber"] != ours[k]["CertNumber"]]
    print(f"\ncert mismatch samples ({len(mism)}):")
    for k in mism[:8]:
        print("   ref:", ref[k]["CertNumber"], "| ours:", ours[k]["CertNumber"])
    tv = [k for k in agreed if ref[k]["TNVED"] and ours[k]["TNVED"]
          and ref[k]["TNVED"] != ours[k]["TNVED"]]
    print(f"\ntnved mismatch samples ({len(tv)}):")
    for k in tv[:8]:
        print("   ref:", ref[k]["TNVED"], "| ours:", ours[k]["TNVED"])


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "validate":
        validate()
    else:
        main()
