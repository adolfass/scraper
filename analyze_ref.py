#!/usr/bin/env python3
"""Offline analysis of Готовый.xlsx to pin down output rules."""
import json, re, sys
from collections import Counter, defaultdict
import openpyxl

wb = openpyxl.load_workbook("Готовый.xlsx", read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = ws.iter_rows(values_only=True)
hdr = list(next(rows))
print("header:", hdr)
idx = {h: i for i, h in enumerate(hdr)}

cert_rows = []          # (code, brand, certnumber, url, tnved, certtype)
n = 0
for r in rows:
    n += 1
    cn = r[idx["CertNumber"]] if "CertNumber" in idx else None
    if cn:
        cert_rows.append((
            str(r[idx["Code"]] or "").strip(),
            str(r[idx["BrandName"]] or "").strip(),
            str(cn).strip(),
            str(r[idx["CertUrl"]] or "").strip(),
            "" if r[idx["TNVED"]] is None else str(r[idx["TNVED"]]).strip(),
            "" if r[idx["CertType"]] is None else str(r[idx["CertType"]]).strip(),
        ))
print(f"total rows: {n}, with cert: {len(cert_rows)}")

# 1. CertType values
print("\nCertType:", Counter(t for *_, t in cert_rows).most_common())

# 2. URL hosts
hosts = Counter(re.sub(r"^https?://([^/]+).*$", r"\1", u) for _, _, _, u, _, _ in cert_rows)
print("hosts:", hosts.most_common())

# 3. TNVED format
tv = Counter(len(t) for *_, t, _ in cert_rows if t)
print("tnved lengths:", tv.most_common(10))
print("tnved samples:", list({t for *_, t, _ in cert_rows if t})[:15])
multi = [t for *_, t, _ in cert_rows if t and not re.fullmatch(r"\d+", t)]
print("non-plain-digit tnved:", len(multi), multi[:10])

# 4. per cert: TNVED constant?
per_cert_tnved = defaultdict(set)
per_cert_urls = defaultdict(set)
per_cert_codes = defaultdict(set)
per_cert_type = defaultdict(set)
for code, brand, cn, url, tv, ct in cert_rows:
    per_cert_tnved[cn].add(tv)
    per_cert_urls[cn].add(url)
    per_cert_codes[cn].add(code.lower())
    per_cert_type[cn].add(ct)
vary = {c: s for c, s in per_cert_tnved.items() if len(s) > (1 if "" in s else 0) and len(s - {""}) > 1}
print(f"\nunique certs: {len(per_cert_tnved)}; certs with >1 distinct non-empty TNVED: {len(vary)}")
for c, s in list(vary.items())[:10]:
    print("  ", c, sorted(s))
print("certs with >1 url:", sum(1 for s in per_cert_urls.values() if len(s) > 1))
print("certs with >1 type:", sum(1 for s in per_cert_type.values() if len(s) > 1))
mixed = {c: s for c, s in per_cert_type.items() if len(s) > 1}
for c, s in list(mixed.items())[:5]:
    print("  mixed type:", c, s)

# 5. codes per cert / rows per cert
sizes = sorted((len(v) for v in per_cert_codes.values()), reverse=True)
print("\ncodes per cert: max", sizes[0], "median", sizes[len(sizes) // 2], "min", sizes[-1])
print("top certs by code count:")
for c, v in sorted(per_cert_codes.items(), key=lambda kv: -len(kv[1]))[:8]:
    print(f"   {len(v):5d}  {c}  {sorted(per_cert_tnved[c])[:2]}")

# 6. TNVED fill rate by source
fsa = [x for x in cert_rows if "pub.fsa.gov.ru" in x[3]]
other = [x for x in cert_rows if "pub.fsa.gov.ru" not in x[3]]
print(f"\nFSA rows {len(fsa)} tnved-filled {sum(1 for x in fsa if x[4])}")
print(f"other rows {len(other)} tnved-filled {sum(1 for x in other if x[4])}")

# 7. how many distinct codes overall have a cert
codes_with_cert = {x[0].lower() for x in cert_rows}
print(f"\ndistinct codes with cert: {len(codes_with_cert)}")

json.dump({"codes_with_cert": sorted(codes_with_cert)},
          open("state/ref_codes_with_cert.json", "w"), ensure_ascii=False)
wb.close()
print("done")
