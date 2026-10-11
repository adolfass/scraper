#!/usr/bin/env python3
"""Brand-level certificates x product-type match.

For every unmatched row we look for a BRAND-level certificate of the same brand
whose product-type text (identification name + TU document name) shares an
automotive product-type stem with the row's ProdName.  That is the defensible
version of "one TU/URL covers several parts": same brand + same product family.
"""
import json
import re
from collections import Counter, defaultdict

import openpyxl

import build_output as B

STEMS = [
    "фильтр", "масл", "воздуш", "топлив", "салон", "колодк", "диск", "тормоз",
    "амортизат", "стоек", "стойк", "пружин", "подшипник", "ремн", "ролик",
    "натяж", "свеч", "катушк", "стартер", "генератор", "насос", "радиатор",
    "термост", "патрубок", "прокладк", "сальник", "втулк", "шаров", "наконечник",
    "рулев", "тяга", "тяг", "шрус", "привод", "вал", "моторчик", "датчик",
    "лямбда", "клапан", "форсунк", "сцеплен", "корзин", "выжимн", "маховик",
    "бампер", "крыл", "фар", "фонар", "зеркал", "стекл", "дворник", "щетк",
    "компрессор", "осушит", "конденсор", "испарит", "вентилятор", "крышк",
    "бачок", "бак", "шланг", "труб", "глушит", "коллектор", "турбин",
    "интеркул", "шарнир", "опор", "сайлентблок", "рессор", "ступиц", "цапф",
    "крестовин", "кардан", "гидравл", "блок", "провод", "реле", "провода",
    "наконечник", "бензонасос", "стеклоподъем", "замок", "щетка", "поддон",
    "коленвал", "распредвал", "поршень", "кольц", "гильз", "турбин",
    "свеча", "накал", "эбу", "тросик", "трос", "печк", "отопит", "кран",
]


def stems(text):
    t = B.norm_text(text)
    return {s for s in STEMS if s in t}


def brand_text(d):
    parts = []
    prod = d.get("product") or {}
    parts += [str(prod.get("fullName") or ""), str(prod.get("name") or "")]
    for idt in prod.get("identifications") or []:
        parts += [str(idt.get("name") or "")]
        tm = idt.get("tradeMark")
        if isinstance(tm, list):
            parts += [str(x) for x in tm]
        elif tm:
            parts.append(str(tm))
    parts.append(str(d.get("tradeMark") or ""))
    return " ".join(parts)


def product_text(d):
    parts = []
    prod = d.get("product") or {}
    parts += [str(prod.get("fullName") or ""), str(prod.get("name") or "")]
    for idt in prod.get("identifications") or []:
        parts += [str(idt.get("name") or ""), str(idt.get("description") or "")]
        for doc in idt.get("documents") or []:
            if isinstance(doc, dict):
                parts.append(str(doc.get("name") or ""))
    return " ".join(parts)


CODEISH = re.compile(r"\b[0-9a-z]{4,}\b")


def is_brand_level(d):
    prod = d.get("product") or {}
    articles = [str(i.get("article") or "").strip()
                for i in prod.get("identifications") or []]
    articles = [a for a in articles if a]
    if not articles:
        return True
    for a in articles:
        if len(CODEISH.findall(a.lower())) >= 3 and re.search(r"\d", a):
            return False
    return True


def main():
    wb = openpyxl.load_workbook("Итоговый ИП_Результат.xlsx",
                                read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    ci, bi, pi, ni = (h.index("Code"), h.index("BrandName"),
                      h.index("ProdName"), h.index("CertNumber"))
    un_rows = []
    for r in it:
        if not (r[ni] and str(r[ni]).strip()):
            un_rows.append((str(r[bi] or "").strip().lower(),
                            str(r[ci] or "").strip(),
                            str(r[pi] or "").strip()))
    wb.close()
    brands = sorted({b for b, _, _ in un_rows if b})

    # index brand-level cards by brand
    bl_by_brand = defaultdict(list)
    cards = 0
    for line in open("certs.jsonl", encoding="utf-8"):
        rec = json.loads(line)
        d = rec.get("data") or {}
        if not d:
            continue
        cards += 1
        if not is_brand_level(d):
            continue
        nb = B.norm_text(brand_text(d))
        pts = stems(product_text(d))
        if not pts:
            continue
        for b in brands:
            if B.brand_matches(b, nb):
                bl_by_brand[b].append((rec.get("id"), d.get("docId") or "",
                                       sorted(pts)))
    print(f"cards={cards} brands with brand-level+type certs={len(bl_by_brand)}")

    matched = []
    for b, code, pname in un_rows:
        ps = stems(pname)
        if not ps:
            continue
        for rid, num, pts in bl_by_brand.get(b, []):
            common = ps & set(pts)
            if common:
                matched.append((b, code, pname, rid, num, sorted(common)))
                break

    print(f"unmatched rows={len(un_rows)}  "
          f"rows with brand-level product-type match: {len(matched)}")
    byb = Counter(m[0] for m in matched)
    print("\nby brand:")
    for b, n in byb.most_common(30):
        print(f"   {n:4d}  {b!r}")
    print("\nsamples:")
    for m in matched[:20]:
        print(f"   {m[0]:18} {m[1]:14} {m[2][:34]:34} -> rid={m[3]} "
              f"({','.join(m[5])})")
    json.dump(matched, open("/tmp/brand_level_match.json", "w"), ensure_ascii=False)


if __name__ == "__main__":
    main()