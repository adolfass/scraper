#!/usr/bin/env python3
"""Carry ref CertUrl for rows where ref HAS a link but our scan found nothing.

Why they were empty: those codes do not appear in any certificate's
productIdentificationArticle in FSA (direct article search returns total=0),
so a code-based scan can never reproduce ref's relation - the ref author
matched by product names / models / manually.  The LINKS themselves are live.

Verification per unique ref URL (all pub.fsa.gov.ru):
  - search certificates (or declarations for /rds/ URLs) by the ref number
    using num2url's fragment+homoglyph queries;
  - require an exact normalized number match;
  - if the live id equals ref's URL id      -> keep ref URL (verified);
  - if it matches under a different id      -> rebuild URL from live id;
  - no match                                -> skip (rows stay empty).

Emits ref_carry_certs.json (same shape as other_certs.json, keyed
"ProdNum|Code"); build_output merges it with existing keys winning.
Gentle: DELAY seconds between requests (default 1.0).
"""
import json, os, time
from collections import defaultdict

import openpyxl
from num2url import clean_number, find_cert

ROOT = os.path.dirname(os.path.abspath(__file__))
REF = os.path.join(ROOT, "Готовый.xlsx")
OUT = os.path.join(ROOT, "ref_carry_certs.json")
DELAY = float(os.environ.get("DELAY", "1.0"))
RSS_URL = "https://pub.fsa.gov.ru/rss/certificate/view/{id}/product"
RDS_URL = "https://pub.fsa.gov.ru/rds/declaration/view/{id}/common"


def load_ref_linked():
    """Rows with CertUrl present -> {url: [(key, number, tnved, type), ...]}"""
    wb = openpyxl.load_workbook(REF, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    idx = {h: i for i, h in enumerate(hdr)}
    by_url = defaultdict(list)
    for r in it:
        url = str(r[idx["CertUrl"]] or "").strip()
        num = str(r[idx["CertNumber"]] or "").strip()
        if not url or not num:
            continue
        key = (f"{str(r[idx['ProdNum']] or '').strip()}|"
               f"{str(r[idx['Code']] or '').strip().lower()}")
        by_url[url].append((key, num,
                            str(r[idx["TNVED"]] or "").strip(),
                            str(r[idx["CertType"]] or "").strip()))
    wb.close()
    return by_url


def verify_url(url, number):
    """-> (verified, final_url, note)"""
    is_decl = "/rds/" in url
    live_id = url.rstrip("/").split("/view/")[-1].split("/")[0] \
        if "/view/" in url else None
    q = clean_number(number)
    best, err = find_cert(q, is_decl, [])
    if err:
        return False, url, f"search err: {err}"
    if best is None:
        return False, url, "number not found live"
    new_id = str(best.get("id"))
    if live_id and new_id == live_id:
        return True, url, f"same id, status={best.get('idStatus')}"
    tmpl = RDS_URL if is_decl else RSS_URL
    return True, tmpl.format(id=new_id), \
        f"live id {new_id} (ref {live_id}), status={best.get('idStatus')}"


def main():
    by_url = load_ref_linked()
    print(f"ref linked urls: {len(by_url)}", flush=True)

    entries = {}
    stats = {"verified": 0, "rebuilt": 0, "miss": 0, "error": 0,
             "rows": 0, "rows_missed": 0}
    for i, (url, keys) in enumerate(sorted(by_url.items(),
                                           key=lambda x: -len(x[1])), 1):
        number = keys[0][1]
        try:
            ok, final_url, note = verify_url(url, number)
        except Exception as e:
            ok, final_url, note = False, url, f"EXC {str(e)[:60]}"
        if ok:
            stats["verified" if final_url == url else "rebuilt"] += 1
            for key, num, tnved, ctype in keys:
                entries[key] = {"number": num, "url": final_url,
                                "tnved": tnved, "type": ctype,
                                "verified": True, "host": "ref_url"}
            stats["rows"] += len(keys)
            print(f"  [{i}/{len(by_url)}] OK ({len(keys)} rows) {note}",
                  flush=True)
        else:
            if "EXC" in note or "err" in note:
                stats["error"] += 1
            else:
                stats["miss"] += 1
            stats["rows_missed"] += len(keys)
            print(f"  [{i}/{len(by_url)}] MISS ({len(keys)} rows): {note}",
                  flush=True)
        time.sleep(DELAY)

    print(f"done: {stats}", flush=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(entries, f, ensure_ascii=False, indent=1)
    print(f"written {OUT}: {len(entries)} entries", flush=True)


if __name__ == "__main__":
    main()
