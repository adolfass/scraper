#!/usr/bin/env python3
"""Match still-unmatched price codes against harvested ANNEX article lists.

Reads the print-form docx exports (which contain the real annex article lists)
from /tmp/annex/docx and /tmp/fsa_docx, extracts article tokens, and finds the
best certificate for each still-unfilled row.

Output: /tmp/annex/fills.json   code -> {cid, brand, brand_ok, cands|[...]}
"""
import io
import json
import os
import re
import zipfile
from collections import defaultdict

import openpyxl

import build_output as B

DOCX_DIRS = ["/tmp/annex/docx", "/tmp/fsa_docx"]
IDX = "/tmp/annex/certs_index.json"
OUT = "Итоговый ИП_Результат.xlsx"
ALNUM = re.compile(r"[^0-9a-zа-яё]")
# Article items are comma/semicolon/whitespace delimited; keeping separators like
# '-' or '.' inside an item (e.g. "19-029313") so whole-item matching is exact.
ITEM_SPLIT = re.compile(r"[,;\s]+")


def load_unmatched():
    wb = openpyxl.load_workbook(OUT, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    ci, bi, ni = h.index("Code"), h.index("BrandName"), h.index("CertNumber")
    un = {}
    for r in it:
        c = str(r[ci] or "").strip().lower()
        if c and not (r[ni] and str(r[ni]).strip()):
            un[c] = str(r[bi] or "").strip().lower()
    wb.close()
    return un


def docx_text(path):
    z = zipfile.ZipFile(path)
    xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", xml))


def quality_ok(code, nrm):
    """Reject short/ambiguous tokens (dates, years, counts)."""
    if len(nrm) < 5:
        return False
    if not re.search(r"[a-zа-я]", nrm) and len(nrm) < 6:
        return False
    return True


def main():
    un = load_unmatched()
    print(f"unmatched codes: {len(un)}", flush=True)
    nrm = {c: ALNUM.sub("", c) for c in un}
    usable = {c for c in un if quality_ok(c, nrm[c])}
    print(f"usable (quality filter): {len(usable)}", flush=True)

    idx = {}
    if os.path.exists(IDX):
        idx = json.load(open(IDX, encoding="utf-8")).get("certs", {})

    files = []
    for d in DOCX_DIRS:
        if os.path.isdir(d):
            for f in os.listdir(d):
                if f.endswith(".docx"):
                    files.append((f[:-5], os.path.join(d, f)))
    print(f"export files: {len(files)}", flush=True)

    hits = defaultdict(list)   # code -> [ {cid, brand_ok} ]
    for n, (cid, path) in enumerate(files, 1):
        try:
            text = docx_text(path)
        except Exception:
            continue
        low = text.lower()
        items = {ALNUM.sub("", w) for w in ITEM_SPLIT.split(low) if w}
        nb = B.norm_text(text)
        for c in usable:
            if nrm[c] not in items:
                continue
            brand = un[c]
            b_ok = B.brand_matches(brand, nb) if brand else False
            hits[c].append({"cid": cid, "brand_ok": b_ok})
        if n % 500 == 0:
            print(f"  scanned {n}/{len(files)} codes_hit={len(hits)}", flush=True)

    # choose best cert per code
    fills = {}
    for c, cands in hits.items():
        def k(h):
            s = idx.get(str(h["cid"]), {})
            return (1 if h["brand_ok"] else 0,
                    1 if s.get("idStatus") == 6 else 0,
                    str(s.get("date") or ""))
        best = max(cands, key=k)
        fills[c] = {"cid": best["cid"], "brand": un[c],
                    "brand_ok": best["brand_ok"],
                    "n_cands": len(cands),
                    "cands": sorted({h["cid"] for h in cands})}

    json.dump(fills, open("/tmp/annex/fills.json", "w"), ensure_ascii=False)
    gated = sum(1 for v in fills.values() if v["brand_ok"])
    print(f"\ncodes matched: {len(fills)}  brand-gated: {gated}  "
          f"no-gate: {len(fills)-gated}", flush=True)
    for c, v in list(fills.items())[:40]:
        print(f"   {c:16} {v['brand']:16} brand_ok={int(v['brand_ok'])} "
              f"cert={v['cid']} cands={v['n_cands']}", flush=True)


if __name__ == "__main__":
    main()
