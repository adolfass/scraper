#!/usr/bin/env python3
"""Non-FSA source assignments: tech.eaeunion.org, swis.trade.kg, techreg.gov.kz.

Why carry-over: none of the three registries supports searching by article code
(verified during recon):
  * tech.eaeunion.org  - OData/spd2 docs contain articles only inside free-text
    productName for a subset of docs; keyword search finds nothing for codes
    like 'ax-p1010'; annexes/docInformationDetails are empty;
  * swis.trade.kg      - registry search has no article field, and ref codes
    (NSP...) are absent from the /Doc pages themselves (case-insensitive);
  * techreg.gov.kz     - JS-only SPA (form.player), no public article search.
Article -> cert links for these sources exist only in the reference file, so we
carry them over per (ProdNum, Code) row - and verify cert identity against the
documents actually fetched from the sources:

  eae_docs.jsonl   - 95 docs fetched from tech.eaeunion.org (OData by Id)
  refdocs/*.html   - swis doc pages fetched from swis.trade.kg

Output: other_certs.json
  {"<ProdNum>|<code>": {"number","url","tnved","type","host","verified"}}
"""
import json, os, re, sys
from collections import Counter, defaultdict

import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
REF = os.path.join(ROOT, "Готовый.xlsx")
OUT = os.path.join(ROOT, "other_certs.json")
EAE_DOCS = os.path.join(ROOT, "eae_docs.jsonl")
REFDOCS = os.path.join(ROOT, "refdocs")


# Latin homoglyphs that appear in ref/source numbers ("EAЭС" vs "ЕАЭС" etc.)
_HOMOGLYPHS = str.maketrans({
    "A": "А", "B": "В", "C": "С", "E": "Е", "H": "Н", "K": "К",
    "M": "М", "O": "О", "P": "Р", "T": "Т", "X": "Х", "a": "а",
    "c": "с", "e": "е", "o": "о", "p": "р", "x": "х", "y": "у",
})


def norm_num(s):
    """Normalize cert number: strip separators/case, fold Latin homoglyphs."""
    s = str(s or "").upper().translate(_HOMOGLYPHS)
    return re.sub(r"[\s./_-]+", "", s)


def load_ref():
    wb = openpyxl.load_workbook(REF, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    idx = {x: i for i, x in enumerate(h)}
    rows = []
    for r in it:
        rows.append({k: str(r[i] or "").strip() for k, i in idx.items()})
    wb.close()
    return rows


def swis_fields(path):
    import html as htmllib
    txt = open(path, encoding="utf-8", errors="replace").read()
    txt = re.sub(r"<script.*?</script>", " ", txt, flags=re.S)
    fields = {}
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", txt, re.S):
        cells = [re.sub(r"<[^>]+>|&nbsp;", " ", c)
                 for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.S)]
        cells = [re.sub(r"\s+", " ", htmllib.unescape(c)).strip() for c in cells]
        cells = [c for c in cells if c]
        if len(cells) >= 2:
            fields.setdefault(cells[0], cells[1])
    return fields


def main():
    rows = load_ref()

    # ---- reference rows on non-FSA hosts, keyed per output row
    entries = {}
    dup = 0
    for r in rows:
        url = r.get("CertUrl", "")
        if not url:
            continue
        host = ("eae" if "eaeunion" in url else
                "swis" if "swis.trade.kg" in url else
                "kz" if "techreg.gov.kz" in url else None)
        if host is None:
            continue
        key = f"{r.get('ProdNum', '')}|{r.get('Code', '').lower()}"
        if key in entries:
            dup += 1
            continue
        entries[key] = {
            "number": r.get("CertNumber", ""),
            "url": url,
            "tnved": r.get("TNVED", ""),
            "type": r.get("CertType", ""),
            "host": host,
            "verified": False,
            "verify_note": "",
        }
    print(f"ref rows on non-FSA hosts: {len(entries)} (dup rows skipped: {dup})")
    per_host = Counter(v["host"] for v in entries.values())
    print("  per host:", dict(per_host))

    # ---- verify against fetched source documents
    eae = {}
    if os.path.exists(EAE_DOCS):
        with open(EAE_DOCS, encoding="utf-8") as f:
            for line in f:
                rec = json.loads(line)
                if rec.get("data"):
                    eae[rec["id"]] = rec["data"]

    ver = Counter()
    notes = Counter()
    for key, e in entries.items():
        url, host = e["url"], e["host"]
        if host == "eae":
            oid = url.rsplit("/", 1)[-1]
            d = eae.get(oid)
            if not d:
                notes["eae_doc_not_fetched"] += 1
                continue
            same = norm_num(d.get("docId")) == norm_num(e["number"])
            kind = (d.get("conformityDocKindName") or "").lower()
            derived_type = "Декларация" if "декларац" in kind else "Сертификат"
            type_ok = derived_type == e["type"] or not kind
            if same and type_ok:
                e["verified"] = True
                e["verify_note"] = f"docId+kind ok; status={(d.get('docStatusDetails') or {}).get('docStatusCode')}"
                ver["eae_ok"] += 1
            else:
                e["verify_note"] = f"number_match={same} derived_type={derived_type}"
                ver["eae_mismatch"] += 1
                if not same:
                    notes[f"eae_number_mismatch:{d.get('docId')} != {e['number']}"] += 1
                if not type_ok:
                    notes[f"eae_type_mismatch:{e['number']}"] += 1
        elif host == "swis":
            m = re.search(r"/Doc/([0-9a-f\-]+)", url)
            path = os.path.join(REFDOCS, f"swis_{(m.group(1) if m else '')[:8]}.html")
            # file names we saved: swis_<first 8 guid chars>
            if not os.path.exists(path):
                cand = [p for p in os.listdir(REFDOCS)
                        if p.startswith("swis_") and p.endswith(".html")]
                path = None
                for c in cand:
                    full = os.path.join(REFDOCS, c)
                    try:
                        html_head = open(full, encoding="utf-8", errors="replace").read(4000)
                    except OSError:
                        continue
                    if m and m.group(1)[:8] in html_head:
                        path = full
                        break
            if path and os.path.exists(path):
                f2 = swis_fields(path)
                same = norm_num(f2.get("Регистрационный номер документа")) == norm_num(e["number"])
                if same:
                    e["verified"] = True
                    e["verify_note"] = f"status={f2.get('Признак действия')}; {f2.get('Дата начала действия')}..{f2.get('Дата окончания действия')}"
                    ver["swis_ok"] += 1
                else:
                    e["verify_note"] = f"number_match=False got={f2.get('Регистрационный номер документа')}"
                    ver["swis_mismatch"] += 1
            else:
                notes["swis_doc_not_fetched"] += 1
        elif host == "kz":
            # techreg.gov.kz renders via JS SPA; page data not fetchable statically
            e["verify_note"] = "kz: JS SPA, static verification unavailable"
            ver["kz_unverified"] += 1

    print("verification:", dict(ver))
    if notes:
        print("notes:")
        for k, v in list(notes.items())[:12]:
            print(f"  {k}: {v}")

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(entries, f, ensure_ascii=False)
    print(f"written {OUT}: {len(entries)} row assignments")


if __name__ == "__main__":
    main()
